"""Dataset-driven professional report content and PDF/Excel exporters."""

from __future__ import annotations

from html import escape
from io import BytesIO
from typing import Any, Dict, List, Optional, Tuple

import matplotlib

matplotlib.use("Agg")

import pandas as pd
from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.figure import Figure
from openpyxl.styles import Font, PatternFill
from pandas.api import types as ptypes
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Image,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from src.analysis import (
    all_outlier_analysis,
    chart_columns,
    classify_columns,
    column_profile,
    correlation_matrix,
    dataset_summary,
    generate_findings,
    numeric_statistics,
    quality_label,
    quality_score,
    strong_correlations,
    time_series_agg,
    top_categories,
)
from src.config import COLORS

PROFILE_COLUMNS = [
    "Column",
    "Kind",
    "Data Type",
    "Non-Null",
    "Missing",
    "Missing %",
    "Unique",
]


def _display_value(value: Any) -> str:
    if pd.isna(value):
        return "—"
    if isinstance(value, float):
        return f"{value:,.4g}"
    return str(value)


def _pdf_text(value: Any) -> str:
    """Keep report text compatible with ReportLab's built-in Helvetica fonts."""
    return str(value).encode("latin-1", errors="replace").decode("latin-1")


def _profile(df: pd.DataFrame) -> pd.DataFrame:
    profile = column_profile(df)
    return profile if not profile.empty else pd.DataFrame(columns=PROFILE_COLUMNS)


def _date_column(df: pd.DataFrame) -> Tuple[Optional[Any], Optional[pd.Series]]:
    for column in df.columns:
        series = df[column]
        if ptypes.is_datetime64_any_dtype(series):
            parsed = pd.to_datetime(series, errors="coerce")
        elif ptypes.is_object_dtype(series) or ptypes.is_string_dtype(series):
            values = series.dropna()
            if values.empty:
                continue
            parsed_values = pd.to_datetime(values, errors="coerce", format="mixed")
            if parsed_values.notna().mean() < 0.8:
                continue
            parsed = pd.to_datetime(series, errors="coerce", format="mixed")
        else:
            continue
        if parsed.notna().any():
            return column, parsed
    return None, None


def _categorical_summary(df: pd.DataFrame) -> pd.DataFrame:
    records = []
    for column in classify_columns(df)["categorical"]:
        series = df[column]
        counts = series.value_counts(dropna=True)
        if counts.empty:
            most_common, frequency = "No non-null values", 0
        else:
            most_common, frequency = str(counts.index[0]), int(counts.iloc[0])
        records.append(
            {
                "Column": column,
                "Non-Null": int(series.notna().sum()),
                "Missing": int(series.isna().sum()),
                "Unique": int(series.nunique(dropna=True)),
                "Most Common Category": most_common,
                "Frequency": frequency,
            }
        )
    return pd.DataFrame(
        records,
        columns=[
            "Column",
            "Non-Null",
            "Missing",
            "Unique",
            "Most Common Category",
            "Frequency",
        ],
    )


def _outlier_table(df: pd.DataFrame) -> pd.DataFrame:
    records = []
    for column, result in all_outlier_analysis(df).items():
        examples = ", ".join(_display_value(value) for value in result["outlier_values"][:5])
        records.append(
            {
                "Column": column,
                "Q1": result["q1"],
                "Q3": result["q3"],
                "IQR": result["iqr"],
                "Lower Bound": result["lower_bound"],
                "Upper Bound": result["upper_bound"],
                "Potential Outliers": result["count"],
                "Outlier %": round(result["pct"], 2),
                "Example Values": examples or "None",
            }
        )
    return pd.DataFrame(
        records,
        columns=[
            "Column",
            "Q1",
            "Q3",
            "IQR",
            "Lower Bound",
            "Upper Bound",
            "Potential Outliers",
            "Outlier %",
            "Example Values",
        ],
    )


def _make_chart(title: str, draw) -> bytes:
    figure = Figure(figsize=(7.4, 3.7), constrained_layout=True)
    axis = figure.subplots()
    draw(axis)
    axis.set_title(title, color=COLORS["navy_900"], fontsize=12, pad=12)
    axis.grid(axis="y", alpha=0.2)
    image = BytesIO()
    FigureCanvasAgg(figure).print_png(image)
    figure.clear()
    return image.getvalue()


def _report_charts(df: pd.DataFrame) -> List[Dict[str, Any]]:
    charts: List[Dict[str, Any]] = []
    suitable = chart_columns(df)
    for column in suitable["numeric"][:4]:
        values = df[column].dropna()
        if len(values) > 10000:
            values = values.sample(n=10000, random_state=42)
        charts.append(
            {
                "title": f"Distribution of {column}",
                "image": _make_chart(
                    f"Distribution of {column}",
                    lambda axis, values=values, column=column: (
                        axis.hist(values, bins=min(30, max(5, int(len(values) ** 0.5))), color=COLORS["blue_500"]),
                        axis.set_xlabel(column),
                        axis.set_ylabel("Count"),
                    ),
                ),
            }
        )

    for column in suitable["categorical"][:3]:
        top = top_categories(df, column, top_n=10)
        labels = top[column].astype(str).tolist()
        counts = top["value"].tolist()

        def draw_categories(axis, labels=labels, counts=counts):
            axis.barh(labels[::-1], counts[::-1], color=COLORS["blue_500"])
            axis.set_xlabel("Rows")

        charts.append(
            {
                "title": f"Top categories in {column}",
                "image": _make_chart(f"Top categories in {column}", draw_categories),
            }
        )

    date_col, parsed_dates = _date_column(df)
    if date_col is not None and parsed_dates is not None:
        trend_df = df.copy()
        trend_df[date_col] = parsed_dates
        numeric = suitable["numeric"]
        if numeric:
            trend = time_series_agg(trend_df, date_col, numeric[0], freq="M", agg="mean")
            value_name = numeric[0]
            trend = trend.rename(columns={numeric[0]: "value"})
            trend_title = f"Monthly mean {value_name} by {date_col}"
        else:
            trend = (
                trend_df.dropna(subset=[date_col])
                .set_index(date_col)
                .resample("ME")
                .size()
                .rename("value")
                .reset_index()
            )
            trend_title = f"Monthly row count by {date_col}"
        if not trend.empty:
            charts.append(
                {
                    "title": trend_title,
                    "image": _make_chart(
                        trend_title,
                        lambda axis, trend=trend: axis.plot(
                            trend.iloc[:, 0], trend["value"], marker="o", color=COLORS["blue_500"]
                        ),
                    ),
                }
            )

    corr = correlation_matrix(df)
    if corr.shape[1] > 1:
        matrix = corr.to_numpy()
        charts.append(
            {
                "title": "Correlation heatmap",
                "image": _make_chart(
                    "Correlation heatmap",
                    lambda axis, matrix=matrix, columns=list(corr.columns): (
                        axis.imshow(matrix, cmap="coolwarm", vmin=-1, vmax=1, aspect="auto"),
                        axis.set_xticks(range(len(columns)), columns, rotation=45, ha="right"),
                        axis.set_yticks(range(len(columns)), columns),
                    ),
                ),
            }
        )
    return charts


def _findings_and_recommendations(
    df: pd.DataFrame,
    profile: pd.DataFrame,
    summary: Dict[str, float],
    outliers: pd.DataFrame,
) -> Tuple[List[str], List[str]]:
    findings = generate_findings(df)
    groups = classify_columns(df)

    if df.empty or df.shape[1] == 0:
        findings.append("The selected dataset has no observations or no columns to analyse.")
    else:
        for column in groups["numeric"]:
            values = df[column].dropna()
            if not values.empty:
                findings.append(
                    f"'{column}' has a median of {_display_value(values.median())} "
                    f"and a range from {_display_value(values.min())} to {_display_value(values.max())}."
                )
                break
        for column in groups["categorical"]:
            counts = df[column].value_counts(dropna=True)
            if not counts.empty:
                findings.append(
                    f"'{column}' is led by '{counts.index[0]}' "
                    f"({int(counts.iloc[0]):,} rows; {counts.iloc[0] / len(df) * 100:.1f}% of all rows)."
                )
                break
        date_col, dates = _date_column(df)
        if date_col is not None and dates is not None and dates.notna().any():
            findings.append(
                f"'{date_col}' spans {dates.min().date()} to {dates.max().date()} "
                f"across {int(dates.nunique())} distinct date values."
            )
        if not outliers.empty and outliers["Potential Outliers"].sum():
            highest = outliers.sort_values("Potential Outliers", ascending=False).iloc[0]
            findings.append(
                f"'{highest['Column']}' has {int(highest['Potential Outliers']):,} "
                f"potential IQR outlier(s) ({highest['Outlier %']:.1f}% of its non-null values)."
            )

    recommendations: List[str] = []
    missing_columns = profile[profile["Missing"] > 0] if not profile.empty else profile
    if summary["missing_cells"]:
        if not missing_columns.empty:
            most_missing = missing_columns.sort_values("Missing %", ascending=False).iloc[0]
            recommendations.append(
                f"Review missingness in '{most_missing['Column']}' ({most_missing['Missing %']:.1f}% missing); "
                "confirm whether the gaps are meaningful before choosing an imputation or exclusion policy."
            )
    else:
        recommendations.append("Keep a missing-value check in the ingestion workflow; this dataset currently has no missing cells.")
    if summary["duplicate_rows"]:
        recommendations.append(
            f"Validate whether the {int(summary['duplicate_rows']):,} duplicate row(s) are redundant using the business key "
            "before removing any records."
        )
    else:
        recommendations.append("Retain duplicate-row checks in future refreshes; none were found in this selected dataset.")
    if not outliers.empty and outliers["Potential Outliers"].sum():
        columns = ", ".join(
            str(row["Column"])
            for _, row in outliers[outliers["Potential Outliers"] > 0].iterrows()
        )
        recommendations.append(
            f"Investigate potential IQR outliers in {columns} against source records and domain limits. "
            "They are flags only and have not been removed."
        )
    if strong_correlations(df):
        recommendations.append(
            "Treat strong correlations as associations, not causal evidence; verify relationships with domain context "
            "before making decisions."
        )
    if not groups["numeric"]:
        recommendations.append("If quantitative comparisons are needed, confirm whether any fields should be represented as numeric values.")
    if not groups["categorical"]:
        recommendations.append("If segmentation is needed, identify suitable category or label fields in the source data.")
    if not recommendations:
        recommendations.append("No issues were identified by the selected checks; continue validating the data as new records arrive.")
    return findings, recommendations


def build_report(df: pd.DataFrame, scope: str) -> Dict[str, Any]:
    """Build all report sections from the selected dataset without mutating it."""
    summary = dataset_summary(df)
    profile = _profile(df)
    stats = numeric_statistics(df)
    categorical = _categorical_summary(df)
    correlations = correlation_matrix(df)
    outliers = _outlier_table(df)
    findings, recommendations = _findings_and_recommendations(
        df, profile, summary, outliers
    )
    score = quality_score(summary)
    conclusion = (
        f"The selected {scope.lower()} contains {summary['rows']:,} rows and "
        f"{summary['columns']:,} columns. Its data-quality score is {score:.0f}/100 "
        f"({quality_label(score)}). Findings and potential outliers in this report are "
        "descriptive signals for review; no values were changed or removed."
    )
    executive_summary = (
        f"This report analyses the {scope.lower()} with {summary['rows']:,} rows and "
        f"{summary['columns']:,} columns. It contains {summary['missing_cells']:,} missing "
        f"cells ({summary['missing_pct']:.1f}%), {summary['duplicate_rows']:,} duplicate "
        f"rows, and a data-quality score of {score:.0f}/100 ({quality_label(score)})."
    )
    return {
        "scope": scope,
        "summary": summary,
        "quality_score": score,
        "quality_label": quality_label(score),
        "profile": profile,
        "numeric_stats": stats,
        "categorical_summary": categorical,
        "correlations": correlations,
        "strong_correlations": strong_correlations(df),
        "outliers": outliers,
        "findings": findings,
        "recommendations": recommendations,
        "executive_summary": executive_summary,
        "conclusion": conclusion,
        "charts": _report_charts(df),
    }


def _as_table_frame(frame: pd.DataFrame, message: str) -> pd.DataFrame:
    if not frame.empty:
        return frame
    return pd.DataFrame({"Information": [message]})


def build_excel_report(df: pd.DataFrame, report: Dict[str, Any]) -> bytes:
    """Export the report into five formatted worksheets."""
    summary = report["summary"]
    overview = pd.DataFrame(
        [
            ("Dataset scope", report["scope"]),
            ("Rows", summary["rows"]),
            ("Columns", summary["columns"]),
            ("Numeric columns", summary["numeric_count"]),
            ("Categorical columns", summary["categorical_count"]),
            ("Date/time columns", summary["datetime_count"]),
            ("Missing cells", summary["missing_cells"]),
            ("Missing cells (%)", round(summary["missing_pct"], 2)),
            ("Duplicate rows", summary["duplicate_rows"]),
            ("Duplicate rows (%)", round(summary["duplicate_pct"], 2)),
            ("Complete rows", summary["complete_rows"]),
            ("Data quality score", round(report["quality_score"], 1)),
            ("Quality label", report["quality_label"]),
            ("Memory usage (bytes)", summary["memory_bytes"]),
        ],
        columns=["Metric", "Value"],
    )
    quality = report["profile"].copy()
    if quality.empty:
        quality = _as_table_frame(quality, "No columns are available for a per-column quality assessment.")
    stats = _as_table_frame(
        report["numeric_stats"], "No numeric columns are available for a statistical summary."
    )
    finding_records = [
        {"Type": "Finding", "Details": finding} for finding in report["findings"]
    ] + [
        {"Type": "Recommendation", "Details": recommendation}
        for recommendation in report["recommendations"]
    ]
    finding_records.append({"Type": "Conclusion", "Details": report["conclusion"]})
    findings = pd.DataFrame(finding_records, columns=["Type", "Details"])
    cleaned = _as_table_frame(df, "The selected dataset contains no rows or columns.")

    workbook = BytesIO()
    sheet_frames = {
        "Dataset Overview": overview,
        "Data Quality": quality,
        "Statistical Summary": stats,
        "Key Findings": findings,
        "Cleaned Dataset": cleaned,
    }
    with pd.ExcelWriter(workbook, engine="openpyxl") as writer:
        for sheet, frame in sheet_frames.items():
            frame.to_excel(writer, index=False, sheet_name=sheet)
        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.fill = PatternFill(fill_type="solid", fgColor="16254F")
                cell.font = Font(bold=True, color="FFFFFF")
            sample_rows = min(worksheet.max_row, 201)
            for cells in worksheet.iter_cols(min_row=1, max_row=sample_rows):
                letter = cells[0].column_letter
                width = min(
                    60,
                    max(
                        12,
                        max(
                            len(str(cell.value)) if cell.value is not None else 0
                            for cell in cells
                        )
                        + 2,
                    ),
                )
                worksheet.column_dimensions[letter].width = width
    return workbook.getvalue()


def _frame_table(frame: pd.DataFrame, available_width: float) -> Table:
    header_style = ParagraphStyle(
        "TableHeader", fontName="Helvetica-Bold", fontSize=7, leading=8, textColor=colors.white
    )
    cell_style = ParagraphStyle(
        "TableCell", fontName="Helvetica", fontSize=7, leading=9, textColor=colors.black
    )
    rows = [[Paragraph(escape(_pdf_text(column)), header_style) for column in frame.columns]]
    for record in frame.itertuples(index=False, name=None):
        rows.append(
            [
                Paragraph(escape(_pdf_text(_display_value(value))), cell_style)
                for value in record
            ]
        )
    widths = [available_width / max(1, len(frame.columns))] * max(1, len(frame.columns))
    table = Table(rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(COLORS["navy_700"])),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("LEADING", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


def build_pdf_report(report: Dict[str, Any]) -> bytes:
    """Render a paginated PDF with report sections, tables and generated charts."""
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=17 * mm,
        bottomMargin=17 * mm,
        title="InsightIQ Professional Data Science Report",
        author="InsightIQ",
    )
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="ReportTitle",
            parent=styles["Title"],
            textColor=colors.HexColor(COLORS["navy_900"]),
            alignment=TA_LEFT,
            fontSize=21,
            leading=25,
            spaceAfter=9,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ReportHeading",
            parent=styles["Heading2"],
            textColor=colors.HexColor(COLORS["blue_500"]),
            fontSize=13,
            leading=16,
            spaceBefore=11,
            spaceAfter=6,
        )
    )
    styles.add(ParagraphStyle(name="ReportBody", parent=styles["BodyText"], fontSize=9, leading=13))
    available_width = document.width
    story = [
        Paragraph("InsightIQ Professional Data Science Report", styles["ReportTitle"]),
        Paragraph(f"Dataset scope: {_pdf_text(report['scope'])}", styles["ReportBody"]),
        Spacer(1, 8),
    ]

    def heading(text: str) -> None:
        story.append(Paragraph(escape(text), styles["ReportHeading"]))

    def paragraph(text: str) -> None:
        story.append(Paragraph(escape(_pdf_text(text)), styles["ReportBody"]))
        story.append(Spacer(1, 4))

    def table(frame: pd.DataFrame, message: str) -> None:
        shown = _as_table_frame(frame, message)
        story.append(_frame_table(shown, available_width))
        story.append(Spacer(1, 8))

    summary = report["summary"]
    heading("Executive Summary")
    paragraph(report["executive_summary"])

    heading("Dataset Overview")
    table(
        pd.DataFrame(
            [
                ("Rows", summary["rows"]),
                ("Columns", summary["columns"]),
                ("Numeric columns", summary["numeric_count"]),
                ("Categorical columns", summary["categorical_count"]),
                ("Date/time columns", summary["datetime_count"]),
                ("Memory usage (bytes)", summary["memory_bytes"]),
                ("Data quality score", f"{report['quality_score']:.0f}/100 — {report['quality_label']}"),
            ],
            columns=["Metric", "Value"],
        ),
        "No dataset metadata is available.",
    )

    heading("Column Names and Data Types")
    table(
        report["profile"][["Column", "Kind", "Data Type", "Non-Null", "Unique"]]
        if not report["profile"].empty
        else pd.DataFrame(),
        "No columns are available in this dataset.",
    )

    heading("Data Quality Assessment")
    table(
        pd.DataFrame(
            [
                ("Missing cells", summary["missing_cells"]),
                ("Missing cells (%)", f"{summary['missing_pct']:.2f}%"),
                ("Duplicate rows", summary["duplicate_rows"]),
                ("Duplicate rows (%)", f"{summary['duplicate_pct']:.2f}%"),
                ("Complete rows", summary["complete_rows"]),
            ],
            columns=["Measure", "Result"],
        ),
        "No quality metrics are available.",
    )

    heading("Missing Values Analysis")
    table(
        report["profile"][["Column", "Missing", "Missing %"]]
        if not report["profile"].empty
        else pd.DataFrame(),
        "No columns are available to assess missing values.",
    )

    heading("Duplicate Rows Analysis")
    paragraph(
        f"{summary['duplicate_rows']:,} duplicate rows were detected "
        f"({summary['duplicate_pct']:.2f}% of rows). Duplicate analysis is informational; "
        "this report does not remove or change any records."
    )

    heading("Statistical Summary for Numeric Columns")
    table(report["numeric_stats"], "No numeric columns are available for statistical summaries.")

    heading("Categorical Column Summaries")
    table(report["categorical_summary"], "No categorical columns are available for summaries.")

    heading("Correlation Analysis")
    if report["strong_correlations"]:
        table(
            pd.DataFrame(
                [
                    {
                        "Column A": first,
                        "Column B": second,
                        "Pearson r": round(value, 4),
                        "Strength": "Strong (|r| >= 0.7)",
                    }
                    for first, second, value in report["strong_correlations"]
                ]
            ),
            "No strong correlations were detected.",
        )
    else:
        paragraph("No strong correlations (absolute Pearson r >= 0.7) were detected.")
    if not report["correlations"].empty and report["correlations"].shape[1] <= 8:
        table(report["correlations"].reset_index(names="Column"), "Correlation matrix is unavailable.")
    else:
        if report["correlations"].empty:
            paragraph("A correlation matrix requires at least two numeric columns with sufficient data.")
        else:
            paragraph(
                "The numeric correlation matrix is included as a heatmap; the PDF limits the detailed "
                "matrix table to eight columns for readability."
            )

    heading("Outlier Analysis (IQR Method)")
    paragraph(
        "Values outside Q1 - 1.5 x IQR or Q3 + 1.5 x IQR are labeled potential outliers. "
        "These are review flags, not proof of errors, and were not automatically deleted."
    )
    table(report["outliers"], "No numeric columns have sufficient data for IQR outlier analysis.")

    heading("Key Data Findings")
    for finding in report["findings"]:
        paragraph(f"- {finding}")

    heading("Practical Recommendations")
    for recommendation in report["recommendations"]:
        paragraph(f"- {recommendation}")

    heading("Conclusion")
    paragraph(report["conclusion"])

    if report["charts"]:
        story.append(PageBreak())
        heading("Charts")
        for chart in report["charts"]:
            story.append(Paragraph(escape(_pdf_text(chart["title"])), styles["ReportBody"]))
            image = Image(BytesIO(chart["image"]))
            image.drawWidth = min(170 * mm, available_width)
            image.drawHeight = image.drawWidth * 0.5
            story.extend([image, Spacer(1, 8)])

    def draw_page_number(canvas, _document) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawRightString(
            landscape(A4)[0] - 15 * mm,
            9 * mm,
            f"InsightIQ Professional Report | Page {_document.page}",
        )
        canvas.restoreState()

    document.build(story, onFirstPage=draw_page_number, onLaterPages=draw_page_number)
    return output.getvalue()
