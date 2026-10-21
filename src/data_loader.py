"""Read uploaded CSV / Excel files into pandas DataFrames, with clear errors."""

from __future__ import annotations

import re
from io import BytesIO
from typing import List, Optional

import pandas as pd
from pandas.api import types as ptypes

CSV_ENCODINGS = ("utf-8-sig", "cp1252", "latin-1")
CSV_DELIMITERS = (",", ";", "\t", "|")

# Only strings that look like "2024-01-31", "31/01/2024" or "2024-01-31 10:30:00"
# are considered for automatic date conversion.
_DATE_PATTERN = re.compile(
    r"^\s*\d{1,4}[-/.]\d{1,2}[-/.]\d{1,4}(?:[ T]\d{1,2}:\d{2}(?::\d{2})?)?\s*$"
)


class DataLoadError(Exception):
    """Raised with a user-friendly message when a file cannot be used."""


def get_extension(file_name: str) -> str:
    return file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""


def list_excel_sheets(data: bytes) -> List[str]:
    """Return the worksheet names of an .xlsx file."""
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(BytesIO(data), read_only=True)
        names = list(workbook.sheetnames)
        workbook.close()
    except Exception as exc:  # corrupt zip, wrong format, password protected...
        raise DataLoadError(
            "This file could not be opened as an Excel workbook. "
            "It may be corrupted, password-protected or not a real .xlsx file."
        ) from exc
    if not names:
        raise DataLoadError("The Excel workbook does not contain any worksheets.")
    return names


def load_dataset(file_name: str, data: bytes, sheet_name: Optional[str] = None) -> pd.DataFrame:
    """Parse the uploaded bytes into a cleaned DataFrame.

    Raises DataLoadError for unsupported, empty, or unreadable files.
    """
    if not data:
        raise DataLoadError("The uploaded file is empty (0 bytes).")

    extension = get_extension(file_name)
    if extension == "csv":
        df = _read_csv(data)
    elif extension == "xlsx":
        df = _read_excel(data, sheet_name)
    else:
        raise DataLoadError(
            f"Unsupported file type '.{extension or 'unknown'}'. Please upload a .csv or .xlsx file."
        )
    return _clean_frame(df)


# ---- Readers --------------------------------------------------------------
def _detect_delimiter(data: bytes, encoding: str) -> str:
    sample = data[:20000].decode(encoding, errors="ignore")
    head = "\n".join(sample.splitlines()[:5])
    counts = {d: head.count(d) for d in CSV_DELIMITERS}
    best = max(counts, key=counts.get)
    return best if counts[best] > 0 else ","


def _read_csv(data: bytes) -> pd.DataFrame:
    for encoding in CSV_ENCODINGS:
        try:
            return pd.read_csv(
                BytesIO(data), sep=_detect_delimiter(data, encoding), encoding=encoding
            )
        except UnicodeDecodeError:
            continue  # try the next encoding
        except pd.errors.EmptyDataError as exc:
            raise DataLoadError("The CSV file is empty or has no readable columns.") from exc
        except pd.errors.ParserError as exc:
            raise DataLoadError(f"The CSV file is malformed and could not be parsed. Details: {exc}") from exc
        except Exception as exc:
            raise DataLoadError(f"The CSV file could not be read. Details: {exc}") from exc
    raise DataLoadError("The CSV file uses an unsupported text encoding.")


def _read_excel(data: bytes, sheet_name: Optional[str]) -> pd.DataFrame:
    try:
        return pd.read_excel(BytesIO(data), sheet_name=sheet_name or 0, engine="openpyxl")
    except ValueError as exc:
        raise DataLoadError(f"The Excel sheet could not be read. Details: {exc}") from exc
    except Exception as exc:
        raise DataLoadError(
            "The Excel file could not be read. It may be corrupted or password-protected."
        ) from exc


# ---- Cleaning -------------------------------------------------------------
def _clean_frame(df: pd.DataFrame) -> pd.DataFrame:
    """Tidy column names, drop fully empty rows/columns and detect date columns."""
    df = df.copy()
    df.columns = _unique_names([str(c).strip() for c in df.columns])
    df = df.dropna(axis=0, how="all").dropna(axis=1, how="all")

    if df.shape[1] == 0:
        raise DataLoadError("The file does not contain any columns with data.")
    if df.shape[0] == 0:
        raise DataLoadError(
            "The dataset has column headers but no data rows. Please upload a file with at least one row."
        )

    df = df.reset_index(drop=True)
    return _convert_date_columns(df)


def _unique_names(names: List[str]) -> List[str]:
    seen: dict = {}
    result = []
    for position, name in enumerate(names, start=1):
        name = name or f"Column {position}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 0
        result.append(name)
    return result


def _convert_date_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Convert text columns that clearly contain dates to datetime64."""
    for column in df.columns:
        series = df[column]
        if (
            ptypes.is_numeric_dtype(series)
            or ptypes.is_bool_dtype(series)
            or ptypes.is_datetime64_any_dtype(series)
        ):
            continue
        non_null = series.dropna()
        if non_null.empty:
            continue
        sample = non_null.head(200).astype(str)
        if sample.str.match(_DATE_PATTERN).mean() < 0.95:
            continue
        parsed = pd.to_datetime(series, errors="coerce", format="mixed")
        if parsed.notna().sum() >= 0.95 * len(non_null):
            df[column] = parsed
    return df
